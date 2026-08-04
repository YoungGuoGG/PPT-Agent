from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from ppt_agent.skills import SkillLoader
from ppt_agent.storage import AgentStore
from ppt_agent.style_library import StyleImportService, StyleLibrary, parse_design


class StyleLibraryTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.data = Path(self.temp.name)
        self.root = Path(__file__).resolve().parents[2]
        self.store = AgentStore(self.data / "agent.db")
        self.library = StyleLibrary(
            self.root / "style-library",
            self.data / "style-library",
            self.store,
        )

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_every_active_style_is_a_valid_published_design_md(self) -> None:
        styles = self.library.list()

        self.assertGreaterEqual(len(styles), 5)
        for style in styles:
            document = parse_design(Path(style["design_path"]), required_status="published")
            self.assertEqual(document.metadata["id"], style["id"])
            self.assertEqual(len(style["design_hash"]), 64)

    def test_pptx_import_sanitizes_references_and_publishes_only_after_review(self) -> None:
        source = self.data / "source.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
        box.text_frame.text = "ACME CONFIDENTIAL CUSTOMER ROADMAP"
        deck.save(source)
        service = StyleImportService(
            self.data / "style-imports",
            self.root / "vendor" / "ppt_master_runtime",
            self.library,
            self.store,
            SkillLoader(self.root / "skills"),
            None,
        )

        created = service.create(source=str(source), name="Imported Editorial")
        draft = service.run(created["id"])

        self.assertEqual(draft["status"], "review_ready")
        self.assertIsNone(self.library.get("imported-editorial"))
        self.assertTrue(draft["references"])
        reference = service.reference_path(created["id"], draft["references"][0]).read_text(encoding="utf-8")
        self.assertNotIn("ACME", reference)
        self.assertRegex(reference, r"\[(TITLE|BODY|LABEL) 1\]")

        published = service.publish(created["id"])

        self.assertEqual(published["id"], "imported-editorial")
        self.assertEqual(published["status"], "published")
        self.assertIsNotNone(self.library.get(published["id"]))


if __name__ == "__main__":
    unittest.main()
